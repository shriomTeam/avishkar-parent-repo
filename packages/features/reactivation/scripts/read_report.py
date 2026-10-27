"""Reads a downloaded Dormant patients report with openpyxl and prints it as JSON, for the browser tests.

uv run python packages/features/reactivation/scripts/read_report.py report.xlsx
"""

import json
import sys

from openpyxl import load_workbook

workbook = load_workbook(sys.argv[1])
print(
    json.dumps(
        {
            "sheets": [
                {
                    "name": ws.title,
                    "freeze": ws.freeze_panes,
                    "rows": [[c.value for c in row] for row in ws.iter_rows()],
                }
                for ws in workbook.worksheets
            ]
        },
        default=str,
    )
)
