import csv
import io
from collections.abc import Iterator
from dataclasses import dataclass
from typing import Any


@dataclass
class CsvSource:
    content: bytes
    source_kind: str = "csv"

    def rows(self) -> Iterator[dict[str, Any]]:
        text = self.content.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        for row in csv.DictReader(io.StringIO(text), dialect=dialect):
            yield {(k or "").strip(): (v.strip() if isinstance(v, str) else v) for k, v in row.items()}

    def headers(self) -> list[str]:
        first = next(self.rows(), None)
        return list(first) if first else []


@dataclass
class ExcelSource:
    content: bytes
    sheet: str | None = None
    source_kind: str = "excel"

    def rows(self) -> Iterator[dict[str, Any]]:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(self.content), read_only=True, data_only=True)
        ws = wb[self.sheet] if self.sheet else wb.active
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it, [])]
        for values in it:
            if all(v is None for v in values):
                continue
            yield {h: v for h, v in zip(header, values, strict=False) if h}

    def headers(self) -> list[str]:
        first = next(self.rows(), None)
        return list(first) if first else []
